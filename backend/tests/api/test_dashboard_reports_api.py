from __future__ import annotations

from datetime import datetime, timezone
from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy.orm import Session, sessionmaker

from app.db.models.health_check import (
    HealthCheck,
    HealthFindingResult,
    PingResult,
)
from tests.api.test_devices_api import DEVICE_PAYLOAD


def _seed_check(
    client: TestClient,
    session_factory: sessionmaker[Session],
    *,
    name: str = "R1",
    status: str = "Warning",
) -> int:
    payload = {**DEVICE_PAYLOAD, "name": name}
    device_id = client.post("/api/v1/devices", json=payload).json()["id"]
    now = datetime.now(timezone.utc)
    with session_factory() as session:
        check = HealthCheck(
            device_id=device_id,
            execution_status="succeeded",
            overall_status=status,
            started_at=now,
            finished_at=now,
            checked_at=now,
            ssh_ok=True,
            cpu_percent=72.0,
            ram_percent=65.0,
        )
        check.pings = [
            PingResult(
                source_type="device",
                target="192.168.141.128",
                packet_loss_percent=10.0,
                latency_min_ms=10.0,
                latency_avg_ms=20.0,
                latency_max_ms=30.0,
            )
        ]
        check.findings = [
            HealthFindingResult(
                severity="Warning",
                metric="cpu",
                message="CPU warning",
            )
        ]
        session.add(check)
        session.commit()
    return device_id


def test_dashboard_summary_uses_latest_device_status(
    client: TestClient,
    session_factory: sessionmaker[Session],
) -> None:
    _seed_check(client, session_factory)

    response = client.get("/api/v1/dashboard/summary")

    assert response.status_code == 200
    body = response.json()
    assert body["total_devices"] == 1
    assert body["warning_devices"] == 1
    assert body["healthy_devices"] == 0
    assert body["critical_devices"] == 0
    assert body["recent_checks"][0]["device_name"] == "R1"


def test_csv_and_pdf_reports_are_downloadable(
    client: TestClient,
    session_factory: sessionmaker[Session],
) -> None:
    _seed_check(client, session_factory, name="=R1")

    csv_response = client.get("/api/v1/reports/csv")
    pdf_response = client.get("/api/v1/reports/pdf")
    xlsx_response = client.get("/api/v1/reports/xlsx")

    assert csv_response.status_code == 200
    assert csv_response.headers["content-type"].startswith("text/csv")
    assert csv_response.content.startswith(b"\xef\xbb\xbf")
    assert b"'=R1" in csv_response.content
    csv_text = csv_response.content.decode("utf-8-sig")
    assert "Mã lần quét" in csv_text
    assert "Thời gian kiểm tra (GMT+7)" in csv_text
    assert "Cảnh báo" in csv_text
    assert "attachment" in csv_response.headers["content-disposition"]

    assert pdf_response.status_code == 200
    assert pdf_response.headers["content-type"] == "application/pdf"
    assert pdf_response.content.startswith(b"%PDF")
    assert len(pdf_response.content) > 2_000

    assert xlsx_response.status_code == 200
    assert xlsx_response.headers["content-type"] == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert xlsx_response.content.startswith(b"PK")
    assert len(xlsx_response.content) > 5_000
    workbook = load_workbook(BytesIO(xlsx_response.content), data_only=False)
    assert workbook.sheetnames == [
        "Tổng quan",
        "Lịch sử kiểm tra",
        "Cổng mạng",
        "Cảnh báo",
    ]
    assert workbook["Tổng quan"]["A2"].value == (
        "BÁO CÁO GIÁM SÁT SỨC KHỎE HỆ THỐNG MẠNG"
    )
    assert workbook["Lịch sử kiểm tra"]["A6"].value == 1
    assert isinstance(workbook["Lịch sử kiểm tra"]["B6"].value, datetime)


def test_reports_reject_reversed_date_range(client: TestClient) -> None:
    response = client.get(
        "/api/v1/reports/csv",
        params={
            "date_from": "2026-10-10T12:00:00Z",
            "date_to": "2026-10-09T12:00:00Z",
        },
    )

    assert response.status_code == 422

