from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass
from datetime import datetime, timedelta, timezone
from pathlib import Path
from statistics import fmean

from sqlalchemy import func, select
from sqlalchemy.orm import Session, selectinload

from app.db.models.device import Device
from app.db.models.health_check import HealthCheck


VIETNAM_TZ = timezone(timedelta(hours=7))


class ReportTooLargeError(Exception):
    def __init__(self, count: int, limit: int) -> None:
        self.count = count
        self.limit = limit
        super().__init__(f"Report has {count} checks; limit is {limit}")


@dataclass(frozen=True, slots=True)
class ReportInterfaceRow:
    name: str
    ip_address: str
    status: str
    protocol: str
    input_errors: int
    crc_errors: int
    output_errors: int
    output_drops: int


@dataclass(frozen=True, slots=True)
class ReportFindingRow:
    severity: str
    metric: str
    message: str


@dataclass(frozen=True, slots=True)
class ReportCheckRow:
    check_id: int
    checked_at: datetime
    status: str
    ssh_ok: bool | None
    uptime_seconds: int | None
    cpu_percent: float | None
    ram_percent: float | None
    ping_target: str
    packet_loss_percent: float | None
    latency_avg_ms: float | None
    error: str
    interfaces: tuple[ReportInterfaceRow, ...]
    findings: tuple[ReportFindingRow, ...]

    @property
    def interfaces_up(self) -> int:
        return sum(
            item.status.lower() == "up" and item.protocol.lower() == "up"
            for item in self.interfaces
        )

    @property
    def warning_findings(self) -> int:
        return sum(item.severity == "Warning" for item in self.findings)

    @property
    def critical_findings(self) -> int:
        return sum(item.severity == "Critical" for item in self.findings)


@dataclass(frozen=True, slots=True)
class ReportRow:
    device_id: int
    device_name: str
    host: str
    site: str
    latest_status: str
    checks_count: int
    cpu_avg: float | None
    cpu_max: float | None
    ram_avg: float | None
    ram_max: float | None
    loss_avg: float | None
    loss_max: float | None
    latency_avg: float | None
    latency_max: float | None
    warning_findings: int
    critical_findings: int
    checks: tuple[ReportCheckRow, ...]


class ReportService:
    MAX_CHECKS = 10_000

    def __init__(self, session: Session) -> None:
        self._session = session

    def rows(
        self,
        *,
        date_from: datetime | None = None,
        date_to: datetime | None = None,
        device_id: int | None = None,
        site: str | None = None,
        status: str | None = None,
    ) -> list[ReportRow]:
        filters = [
            Device.deleted_at.is_(None),
            HealthCheck.execution_status == "succeeded",
        ]
        if date_from is not None:
            filters.append(HealthCheck.started_at >= date_from)
        if date_to is not None:
            filters.append(HealthCheck.started_at <= date_to)
        if device_id is not None:
            filters.append(Device.id == device_id)
        if site is not None:
            filters.append(Device.site == site)
        if status is not None:
            filters.append(HealthCheck.overall_status == status)

        check_count = int(
            self._session.scalar(
                select(func.count(HealthCheck.id))
                .select_from(Device)
                .join(HealthCheck, HealthCheck.device_id == Device.id)
                .where(*filters)
            )
            or 0
        )
        if check_count > self.MAX_CHECKS:
            raise ReportTooLargeError(check_count, self.MAX_CHECKS)

        records = self._session.execute(
            select(Device, HealthCheck)
            .join(HealthCheck, HealthCheck.device_id == Device.id)
            .where(*filters)
            .options(
                selectinload(HealthCheck.interfaces),
                selectinload(HealthCheck.pings),
                selectinload(HealthCheck.findings),
            )
            .order_by(Device.id, HealthCheck.started_at, HealthCheck.id)
        ).all()
        grouped: dict[int, tuple[Device, list[HealthCheck]]] = {}
        for device, check in records:
            grouped.setdefault(device.id, (device, []))[1].append(check)

        result: list[ReportRow] = []
        for device, checks in grouped.values():
            cpu_values = [x.cpu_percent for x in checks if x.cpu_percent is not None]
            ram_values = [x.ram_percent for x in checks if x.ram_percent is not None]
            detail_rows = tuple(_check_row(check) for check in checks)
            losses = [
                check.packet_loss_percent
                for check in detail_rows
                if check.packet_loss_percent is not None
            ]
            latencies = [
                check.latency_avg_ms
                for check in detail_rows
                if check.latency_avg_ms is not None
            ]
            result.append(
                ReportRow(
                    device_id=device.id,
                    device_name=device.name,
                    host=device.host,
                    site=device.site or "Chưa khai báo",
                    latest_status=checks[-1].overall_status or "Unknown",
                    checks_count=len(checks),
                    cpu_avg=_average(cpu_values),
                    cpu_max=max(cpu_values, default=None),
                    ram_avg=_average(ram_values),
                    ram_max=max(ram_values, default=None),
                    loss_avg=_average(losses),
                    loss_max=max(losses, default=None),
                    latency_avg=_average(latencies),
                    latency_max=max(latencies, default=None),
                    warning_findings=sum(
                        check.warning_findings for check in detail_rows
                    ),
                    critical_findings=sum(
                        check.critical_findings for check in detail_rows
                    ),
                    checks=detail_rows,
                )
            )
        return result

    @staticmethod
    def build_csv(rows: list[ReportRow]) -> bytes:
        """Export one row per health check instead of hiding history in aggregates."""
        stream = io.StringIO(newline="")
        writer = csv.writer(stream)
        writer.writerow(
            [
                "Mã lần quét",
                "Thời gian kiểm tra (GMT+7)",
                "Thiết bị",
                "IP quản trị",
                "Khu vực",
                "Trạng thái",
                "SSH",
                "Uptime",
                "CPU (%)",
                "RAM (%)",
                "Đích ping",
                "Packet loss (%)",
                "Latency TB (ms)",
                "Cổng up/tổng",
                "Cảnh báo",
                "Nghiêm trọng",
                "Lỗi kết nối",
            ]
        )
        for row in rows:
            for check in row.checks:
                writer.writerow(
                    [
                        check.check_id,
                        _format_datetime(check.checked_at),
                        _safe_csv(row.device_name),
                        row.host,
                        _safe_csv(row.site),
                        _status_label(check.status),
                        _ssh_label(check.ssh_ok),
                        _format_uptime(check.uptime_seconds),
                        _display(check.cpu_percent),
                        _display(check.ram_percent),
                        check.ping_target,
                        _display(check.packet_loss_percent),
                        _display(check.latency_avg_ms),
                        f"{check.interfaces_up}/{len(check.interfaces)}",
                        check.warning_findings,
                        check.critical_findings,
                        _safe_csv(_translate_error(check.error)),
                    ]
                )
        return b"\xef\xbb\xbf" + stream.getvalue().encode("utf-8")

    @staticmethod
    def build_xlsx(
        rows: list[ReportRow],
        *,
        date_from: datetime | None = None,
        date_to: datetime | None = None,
    ) -> bytes:
        """Build a formatted Excel workbook with summary and supporting detail."""
        from openpyxl import Workbook

        workbook = Workbook()
        overview = workbook.active
        overview.title = "Tổng quan"
        history = workbook.create_sheet("Lịch sử kiểm tra")
        interfaces = workbook.create_sheet("Cổng mạng")
        findings = workbook.create_sheet("Cảnh báo")

        all_checks = [check for row in rows for check in row.checks]
        _build_overview_sheet(
            overview,
            rows,
            all_checks,
            date_from=date_from,
            date_to=date_to,
        )
        _build_history_sheet(history, rows)
        _build_interfaces_sheet(interfaces, rows)
        _build_findings_sheet(findings, rows)

        workbook.properties.title = "Báo cáo giám sát sức khỏe hệ thống mạng"
        workbook.properties.subject = "CPU, RAM, SSH, ping, cổng mạng và cảnh báo"
        workbook.properties.creator = "Network Health Check"
        workbook.properties.description = (
            "Báo cáo sức khỏe thiết bị Cisco IOS theo giờ Việt Nam GMT+7"
        )
        workbook.active = 0
        output = io.BytesIO()
        workbook.save(output)
        return output.getvalue()

    @staticmethod
    def build_pdf(
        rows: list[ReportRow],
        *,
        date_from: datetime | None = None,
        date_to: datetime | None = None,
    ) -> bytes:
        from reportlab.lib import colors
        from reportlab.lib.enums import TA_CENTER, TA_LEFT
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import mm
        from reportlab.platypus import (
            LongTable,
            PageBreak,
            Paragraph,
            SimpleDocTemplate,
            Spacer,
            Table,
            TableStyle,
        )

        font_name, bold_font_name = _register_pdf_fonts()
        output = io.BytesIO()
        page_width, page_height = landscape(A4)
        document = SimpleDocTemplate(
            output,
            pagesize=(page_width, page_height),
            leftMargin=10 * mm,
            rightMargin=10 * mm,
            topMargin=14 * mm,
            bottomMargin=14 * mm,
            title="Báo cáo giám sát sức khỏe hệ thống mạng",
            author="Network Health Check",
            subject="CPU, RAM, SSH, ping, interface và cảnh báo thiết bị mạng",
        )
        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            "ReportTitle",
            parent=styles["Title"],
            fontName=bold_font_name,
            fontSize=20,
            leading=25,
            textColor=colors.HexColor("#153E75"),
            alignment=TA_CENTER,
            spaceAfter=3 * mm,
        )
        section_style = ParagraphStyle(
            "SectionTitle",
            parent=styles["Heading2"],
            fontName=bold_font_name,
            fontSize=12,
            leading=15,
            textColor=colors.HexColor("#153E75"),
            spaceBefore=3 * mm,
            spaceAfter=2 * mm,
        )
        normal_style = ParagraphStyle(
            "ReportNormal",
            parent=styles["Normal"],
            fontName=font_name,
            fontSize=8.5,
            leading=11,
            textColor=colors.HexColor("#334155"),
        )
        small_style = ParagraphStyle(
            "ReportSmall",
            parent=normal_style,
            fontSize=7,
            leading=9,
        )
        header_cell_style = ParagraphStyle(
            "HeaderCell",
            parent=small_style,
            fontName=bold_font_name,
            textColor=colors.white,
            alignment=TA_CENTER,
        )
        body_cell_style = ParagraphStyle(
            "BodyCell",
            parent=small_style,
            alignment=TA_LEFT,
        )

        all_checks = [check for row in rows for check in row.checks]
        generated_at = datetime.now(VIETNAM_TZ)
        period_from, period_to = _report_period(all_checks, date_from, date_to)
        status_counts = {
            key: sum(check.status == key for check in all_checks)
            for key in ("Healthy", "Warning", "Critical")
        }
        warning_count = sum(check.warning_findings for check in all_checks)
        critical_count = sum(check.critical_findings for check in all_checks)

        story = [
            Paragraph("BÁO CÁO GIÁM SÁT SỨC KHỎE HỆ THỐNG MẠNG", title_style),
            Table(
                [
                    [
                        Paragraph("Thời điểm xuất", small_style),
                        Paragraph(_format_datetime(generated_at), normal_style),
                        Paragraph("Múi giờ", small_style),
                        Paragraph("Việt Nam (GMT+7)", normal_style),
                    ],
                    [
                        Paragraph("Khoảng dữ liệu", small_style),
                        Paragraph(_period_label(period_from, period_to), normal_style),
                        Paragraph("Phạm vi", small_style),
                        Paragraph(
                            f"{len(rows)} thiết bị - {len(all_checks)} lần kiểm tra",
                            normal_style,
                        ),
                    ],
                ],
                colWidths=[27 * mm, 78 * mm, 22 * mm, 120 * mm],
                style=TableStyle(
                    [
                        ("FONTNAME", (0, 0), (-1, -1), font_name),
                        ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#E8F1FB")),
                        ("BACKGROUND", (2, 0), (2, -1), colors.HexColor("#E8F1FB")),
                        ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
                        ("INNERGRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#E2E8F0")),
                        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                        ("LEFTPADDING", (0, 0), (-1, -1), 5),
                        ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                        ("TOPPADDING", (0, 0), (-1, -1), 4),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                    ]
                ),
            ),
            Spacer(1, 4 * mm),
            Paragraph("1. Tổng quan", section_style),
            _kpi_table(
                [
                    ("Thiết bị", len(rows), "#2563EB"),
                    ("Lần kiểm tra", len(all_checks), "#475569"),
                    ("Tốt", status_counts["Healthy"], "#059669"),
                    ("Cảnh báo", status_counts["Warning"], "#D97706"),
                    ("Nghiêm trọng", status_counts["Critical"], "#DC2626"),
                    ("Phát hiện cảnh báo/nghiêm trọng", f"{warning_count}/{critical_count}", "#7C3AED"),
                ],
                font_name,
                bold_font_name,
            ),
            Spacer(1, 3 * mm),
            Paragraph("2. Tổng hợp theo thiết bị", section_style),
        ]

        summary_headers = [
            "Thiết bị",
            "IP",
            "Khu vực",
            "Trạng thái mới nhất",
            "Số lần",
            "CPU TB/Max (%)",
            "RAM TB/Max (%)",
            "Loss TB/Max (%)",
            "Latency TB/Max (ms)",
            "Cảnh báo/Nghiêm trọng",
        ]
        summary_data = [[Paragraph(item, header_cell_style) for item in summary_headers]]
        for row in rows:
            summary_data.append(
                [
                    _cell(row.device_name, body_cell_style),
                    _cell(row.host, body_cell_style),
                    _cell(row.site, body_cell_style),
                    _cell(_status_label(row.latest_status), body_cell_style),
                    str(row.checks_count),
                    f"{_display(row.cpu_avg)}/{_display(row.cpu_max)}",
                    f"{_display(row.ram_avg)}/{_display(row.ram_max)}",
                    f"{_display(row.loss_avg)}/{_display(row.loss_max)}",
                    f"{_display(row.latency_avg)}/{_display(row.latency_max)}",
                    f"{row.warning_findings}/{row.critical_findings}",
                ]
            )
        if not rows:
            summary_data.append(
                [Paragraph("Không có dữ liệu phù hợp với bộ lọc.", body_cell_style)]
                + [""] * 9
            )
        summary_table = LongTable(
            summary_data,
            repeatRows=1,
            colWidths=[value * mm for value in [25, 27, 28, 30, 15, 27, 27, 27, 32, 17]],
            hAlign="LEFT",
        )
        summary_table.setStyle(_report_table_style(font_name))
        _apply_status_colors(summary_table, rows, status_column=3)
        story.append(summary_table)

        if all_checks:
            story.extend(
                [
                    PageBreak(),
                    Paragraph("3. Chi tiết từng lần kiểm tra", section_style),
                    Paragraph(
                        "Mỗi dòng là một lần quét thực tế. Dấu '-' nghĩa là thiết bị không trả về chỉ số.",
                        normal_style,
                    ),
                    Spacer(1, 2 * mm),
                ]
            )
            detail_headers = [
                "ID",
                "Thời gian (GMT+7)",
                "Thiết bị",
                "Trạng thái",
                "SSH",
                "Uptime",
                "CPU",
                "RAM",
                "Loss",
                "Latency",
                "Cổng Up",
                "Cảnh báo/Nghiêm trọng",
                "Ghi chú",
            ]
            detail_data = [[Paragraph(item, header_cell_style) for item in detail_headers]]
            detail_statuses: list[str] = []
            for row in rows:
                for check in row.checks:
                    detail_statuses.append(check.status)
                    detail_data.append(
                        [
                            str(check.check_id),
                            _cell(_format_datetime(check.checked_at), body_cell_style),
                            _cell(row.device_name, body_cell_style),
                            _cell(_status_label(check.status), body_cell_style),
                            _ssh_label(check.ssh_ok),
                            _cell(_format_uptime(check.uptime_seconds), body_cell_style),
                            _percent(check.cpu_percent),
                            _percent(check.ram_percent),
                            _percent(check.packet_loss_percent),
                            _milliseconds(check.latency_avg_ms),
                            f"{check.interfaces_up}/{len(check.interfaces)}",
                            f"{check.warning_findings}/{check.critical_findings}",
                            _cell(_translate_error(check.error), body_cell_style),
                        ]
                    )
            detail_table = LongTable(
                detail_data,
                repeatRows=1,
                colWidths=[
                    value * mm
                    for value in [11, 32, 23, 23, 18, 25, 16, 16, 17, 20, 18, 14, 44]
                ],
                hAlign="LEFT",
            )
            detail_table.setStyle(_report_table_style(font_name))
            _apply_status_colors(detail_table, detail_statuses, status_column=3)
            story.append(detail_table)

            latest_interfaces = [
                (row, row.checks[-1], interface)
                for row in rows
                if row.checks
                for interface in row.checks[-1].interfaces
            ]
            story.extend(
                [
                    PageBreak(),
                    Paragraph("4. Trạng thái cổng mạng mới nhất", section_style),
                    Paragraph(
                        "Bảng lấy kết quả cổng mạng từ lần kiểm tra mới nhất của mỗi thiết bị trong phạm vi báo cáo.",
                        normal_style,
                    ),
                    Spacer(1, 2 * mm),
                ]
            )
            interface_headers = [
                "Thiết bị",
                "Thời gian",
                "Interface",
                "Địa chỉ IP",
                "Status/Protocol",
                "Input errors",
                "CRC",
                "Output errors",
                "Output drops",
            ]
            interface_data = [
                [Paragraph(item, header_cell_style) for item in interface_headers]
            ]
            for row, check, interface in latest_interfaces:
                interface_data.append(
                    [
                        _cell(row.device_name, body_cell_style),
                        _cell(_format_datetime(check.checked_at), body_cell_style),
                        _cell(interface.name, body_cell_style),
                        _cell(interface.ip_address or "-", body_cell_style),
                        f"{interface.status}/{interface.protocol}",
                        str(interface.input_errors),
                        str(interface.crc_errors),
                        str(interface.output_errors),
                        str(interface.output_drops),
                    ]
                )
            if not latest_interfaces:
                interface_data.append(
                    [Paragraph("Không có dữ liệu cổng mạng.", body_cell_style)]
                    + [""] * 8
                )
            interface_table = LongTable(
                interface_data,
                repeatRows=1,
                colWidths=[value * mm for value in [25, 34, 35, 34, 30, 27, 18, 27, 26]],
                hAlign="LEFT",
            )
            interface_table.setStyle(_report_table_style(font_name))
            story.append(interface_table)

            finding_entries = [
                (row, check, finding)
                for row in rows
                for check in row.checks
                for finding in check.findings
            ]
            story.extend(
                [
                    PageBreak(),
                    Paragraph("5. Cảnh báo và phát hiện", section_style),
                ]
            )
            if finding_entries:
                finding_headers = [
                    "Thời gian (GMT+7)",
                    "Thiết bị",
                    "Mức độ",
                    "Chỉ số",
                    "Nội dung",
                ]
                finding_data = [
                    [Paragraph(item, header_cell_style) for item in finding_headers]
                ]
                finding_statuses: list[str] = []
                for row, check, finding in finding_entries:
                    finding_statuses.append(finding.severity)
                    finding_data.append(
                        [
                            _cell(_format_datetime(check.checked_at), body_cell_style),
                            _cell(row.device_name, body_cell_style),
                            _cell(_status_label(finding.severity), body_cell_style),
                            _cell(_metric_label(finding.metric), body_cell_style),
                            _cell(
                                _translate_finding_message(finding.message),
                                body_cell_style,
                            ),
                        ]
                    )
                finding_table = LongTable(
                    finding_data,
                    repeatRows=1,
                    colWidths=[value * mm for value in [35, 27, 25, 35, 145]],
                    hAlign="LEFT",
                )
                finding_table.setStyle(_report_table_style(font_name))
                _apply_status_colors(
                    finding_table,
                    finding_statuses,
                    status_column=2,
                )
                story.append(finding_table)
            else:
                story.append(
                    Paragraph(
                        "Không ghi nhận cảnh báo hoặc sự cố trong phạm vi báo cáo.",
                        normal_style,
                    )
                )

        def add_page_frame(canvas, doc) -> None:  # noqa: ANN001
            canvas.saveState()
            canvas.setFont(font_name, 7)
            canvas.setFillColor(colors.HexColor("#64748B"))
            canvas.drawString(
                10 * mm,
                page_height - 8 * mm,
                "NETWORK HEALTH CHECK - BÁO CÁO GIÁM SÁT",
            )
            canvas.setStrokeColor(colors.HexColor("#CBD5E1"))
            canvas.line(10 * mm, page_height - 9.5 * mm, page_width - 10 * mm, page_height - 9.5 * mm)
            canvas.line(10 * mm, 10 * mm, page_width - 10 * mm, 10 * mm)
            canvas.drawString(10 * mm, 6 * mm, "Cisco IOS - GNS3 - Netmiko")
            canvas.drawRightString(
                page_width - 10 * mm,
                6 * mm,
                f"Trang {doc.page}",
            )
            canvas.restoreState()

        document.build(
            story,
            onFirstPage=add_page_frame,
            onLaterPages=add_page_frame,
        )
        return output.getvalue()


def _build_overview_sheet(
    sheet,
    rows: list[ReportRow],
    checks: list[ReportCheckRow],
    *,
    date_from: datetime | None,
    date_to: datetime | None,
) -> None:
    from openpyxl.chart import DoughnutChart, Reference
    from openpyxl.chart.label import DataLabelList
    from openpyxl.styles import Alignment, Font, PatternFill

    _prepare_excel_sheet(sheet, "1F4E78")
    sheet.sheet_view.zoomScale = 90
    sheet["A2"] = "BÁO CÁO GIÁM SÁT SỨC KHỎE HỆ THỐNG MẠNG"
    sheet["A2"].font = Font(name="Arial", size=16, bold=True, color="153E75")
    sheet["A3"] = "Dữ liệu CPU, RAM, SSH, ping, cổng mạng và cảnh báo thiết bị Cisco IOS"
    sheet["A3"].font = Font(name="Arial", size=10, italic=True, color="64748B")

    generated_at = datetime.now(VIETNAM_TZ)
    period_from, period_to = _report_period(checks, date_from, date_to)
    metadata = [
        ("A4", "Thời điểm xuất", "B4", _excel_datetime(generated_at)),
        ("D4", "Múi giờ", "E4", "Việt Nam (GMT+7)"),
        ("G4", "Phạm vi", "H4", f"{len(rows)} thiết bị - {len(checks)} lần kiểm tra"),
        ("A5", "Từ thời điểm", "B5", _excel_datetime(period_from)),
        ("D5", "Đến thời điểm", "E5", _excel_datetime(period_to)),
    ]
    label_fill = PatternFill("solid", fgColor="E8F1FB")
    for label_cell, label, value_cell, value in metadata:
        sheet[label_cell] = label
        sheet[label_cell].font = Font(name="Arial", size=10, bold=True, color="334155")
        sheet[label_cell].fill = label_fill
        sheet[label_cell].alignment = Alignment(vertical="center")
        sheet[value_cell] = value
        sheet[value_cell].font = Font(name="Arial", size=10, color="334155")
        if isinstance(value, datetime):
            sheet[value_cell].number_format = "dd/mm/yyyy hh:mm:ss"

    status_counts = {
        key: sum(check.status == key for check in checks)
        for key in ("Healthy", "Warning", "Critical")
    }
    warning_findings = sum(check.warning_findings for check in checks)
    critical_findings = sum(check.critical_findings for check in checks)
    kpis = [
        ("Thiết bị", len(rows), "2563EB"),
        ("Lần kiểm tra", len(checks), "475569"),
        ("Tốt", status_counts["Healthy"], "059669"),
        ("Cảnh báo", status_counts["Warning"], "D97706"),
        ("Nghiêm trọng", status_counts["Critical"], "DC2626"),
        ("Phát hiện cảnh báo", warning_findings, "D97706"),
        ("Phát hiện nghiêm trọng", critical_findings, "DC2626"),
    ]
    for index, (label, value, color) in enumerate(kpis, start=1):
        label_cell = sheet.cell(row=7, column=index, value=label)
        value_cell = sheet.cell(row=8, column=index, value=value)
        label_cell.font = Font(name="Arial", size=9, color="64748B")
        value_cell.font = Font(name="Arial", size=15, bold=True, color=color)
        for cell in (label_cell, value_cell):
            cell.fill = PatternFill("solid", fgColor="F8FAFC")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.row_dimensions[7].height = 28
    sheet.row_dimensions[8].height = 28

    headers = [
        "Thiết bị",
        "IP quản trị",
        "Khu vực",
        "Trạng thái mới nhất",
        "Số lần kiểm tra",
        "CPU TB",
        "CPU Max",
        "RAM TB",
        "RAM Max",
        "Loss TB",
        "Loss Max",
        "Latency TB (ms)",
        "Latency Max (ms)",
        "Cảnh báo",
        "Nghiêm trọng",
    ]
    header_row = 11
    sheet.append([])
    for column, header in enumerate(headers, start=1):
        sheet.cell(row=header_row, column=column, value=header)
    for row_index, row in enumerate(rows, start=header_row + 1):
        values = [
            _safe_excel_text(row.device_name),
            row.host,
            _safe_excel_text(row.site),
            _status_label(row.latest_status),
            row.checks_count,
            _excel_percent(row.cpu_avg),
            _excel_percent(row.cpu_max),
            _excel_percent(row.ram_avg),
            _excel_percent(row.ram_max),
            _excel_percent(row.loss_avg),
            _excel_percent(row.loss_max),
            row.latency_avg,
            row.latency_max,
            row.warning_findings,
            row.critical_findings,
        ]
        for column, value in enumerate(values, start=1):
            sheet.cell(row=row_index, column=column, value=value)
        for column in range(6, 12):
            if column != 12:
                sheet.cell(row=row_index, column=column).number_format = "0.0%"
        sheet.cell(row=row_index, column=12).number_format = "0.0"
        sheet.cell(row=row_index, column=13).number_format = "0.0"
        _set_excel_status_style(sheet.cell(row=row_index, column=4), row.latest_status)

    _style_excel_header(sheet, header_row, len(headers))
    _style_excel_body(sheet, header_row + 1, header_row + len(rows), len(headers))
    if rows:
        _add_excel_table(sheet, "DeviceSummaryTable", header_row, header_row + len(rows), len(headers))
    else:
        sheet.cell(row=header_row + 1, column=1, value="Không có dữ liệu phù hợp với bộ lọc.")

    widths = [18, 18, 22, 22, 16, 12, 12, 12, 12, 12, 12, 17, 17, 13, 15]
    _set_excel_widths(sheet, widths)
    sheet.freeze_panes = "A12"

    chart = DoughnutChart()
    chart.title = "Phân bố trạng thái các lần kiểm tra"
    chart.add_data(
        Reference(sheet, min_col=3, max_col=5, min_row=8, max_row=8),
        from_rows=True,
    )
    chart.set_categories(Reference(sheet, min_col=3, max_col=5, min_row=7, max_row=7))
    chart.height = 7
    chart.width = 10
    chart.holeSize = 52
    chart.firstSliceAng = 270
    chart.dataLabels = DataLabelList()
    chart.dataLabels.showVal = True
    chart.legend.position = "r"
    sheet.add_chart(chart, "Q3")
    _configure_excel_print(sheet, landscape=True, repeat_rows="11:11")


def _build_history_sheet(sheet, rows: list[ReportRow]) -> None:
    from openpyxl.styles import Alignment, Font

    _prepare_detail_sheet(
        sheet,
        "LỊCH SỬ KIỂM TRA CHI TIẾT",
        "Mỗi dòng là một lần quét. Thời gian hiển thị theo GMT+7.",
        "2563EB",
    )
    headers = [
        "Mã lần quét",
        "Thời gian kiểm tra",
        "Thiết bị",
        "IP quản trị",
        "Khu vực",
        "Trạng thái",
        "SSH",
        "Uptime",
        "CPU",
        "RAM",
        "Đích ping",
        "Packet loss",
        "Latency TB (ms)",
        "Cổng Up",
        "Tổng cổng",
        "Cảnh báo",
        "Nghiêm trọng",
        "Lỗi kết nối",
    ]
    header_row = 5
    for column, header in enumerate(headers, start=1):
        sheet.cell(row=header_row, column=column, value=header)
    current_row = header_row
    for device in rows:
        for check in device.checks:
            current_row += 1
            values = [
                check.check_id,
                _excel_datetime(check.checked_at),
                _safe_excel_text(device.device_name),
                device.host,
                _safe_excel_text(device.site),
                _status_label(check.status),
                _ssh_label(check.ssh_ok),
                check.uptime_seconds / 86_400 if check.uptime_seconds is not None else None,
                _excel_percent(check.cpu_percent),
                _excel_percent(check.ram_percent),
                check.ping_target,
                _excel_percent(check.packet_loss_percent),
                check.latency_avg_ms,
                check.interfaces_up,
                len(check.interfaces),
                check.warning_findings,
                check.critical_findings,
                _safe_excel_text(_translate_error(check.error)),
            ]
            for column, value in enumerate(values, start=1):
                sheet.cell(row=current_row, column=column, value=value)
            sheet.cell(row=current_row, column=2).number_format = "dd/mm/yyyy hh:mm:ss"
            sheet.cell(row=current_row, column=8).number_format = '[h]" giờ "mm" phút"'
            for column in (9, 10, 12):
                sheet.cell(row=current_row, column=column).number_format = "0.0%"
            sheet.cell(row=current_row, column=13).number_format = "0.0"
            _set_excel_status_style(sheet.cell(row=current_row, column=6), check.status)
            _set_excel_ssh_style(sheet.cell(row=current_row, column=7), check.ssh_ok)
            sheet.cell(row=current_row, column=18).alignment = Alignment(
                horizontal="left", vertical="top", wrap_text=True
            )

    _style_excel_header(sheet, header_row, len(headers))
    _style_excel_body(sheet, header_row + 1, current_row, len(headers))
    if current_row > header_row:
        _add_excel_table(sheet, "HealthCheckHistoryTable", header_row, current_row, len(headers))
    else:
        sheet.cell(row=header_row + 1, column=1, value="Không có dữ liệu lịch sử kiểm tra.")
    _set_excel_widths(
        sheet,
        [14, 21, 16, 18, 20, 17, 14, 16, 11, 11, 18, 15, 17, 11, 11, 12, 14, 55],
    )
    sheet.freeze_panes = "A6"
    sheet.auto_filter.ref = f"A{header_row}:R{max(current_row, header_row)}"
    sheet["A3"].font = Font(name="Arial", size=10, italic=True, color="64748B")
    _configure_excel_print(sheet, landscape=True, repeat_rows="5:5")


def _build_interfaces_sheet(sheet, rows: list[ReportRow]) -> None:
    _prepare_detail_sheet(
        sheet,
        "TRẠNG THÁI CỔNG MẠNG MỚI NHẤT",
        "Kết quả cổng từ lần kiểm tra mới nhất của mỗi thiết bị trong phạm vi báo cáo.",
        "0EA5E9",
    )
    headers = [
        "Thời gian kiểm tra",
        "Thiết bị",
        "IP quản trị",
        "Interface",
        "Địa chỉ IP",
        "Status",
        "Protocol",
        "Input errors",
        "CRC errors",
        "Output errors",
        "Output drops",
        "Mã lần quét",
    ]
    header_row = 5
    for column, header in enumerate(headers, start=1):
        sheet.cell(row=header_row, column=column, value=header)
    current_row = header_row
    for device in rows:
        if not device.checks:
            continue
        check = device.checks[-1]
        for interface in check.interfaces:
            current_row += 1
            values = [
                _excel_datetime(check.checked_at),
                _safe_excel_text(device.device_name),
                device.host,
                _safe_excel_text(interface.name),
                interface.ip_address or "-",
                interface.status,
                interface.protocol,
                interface.input_errors,
                interface.crc_errors,
                interface.output_errors,
                interface.output_drops,
                check.check_id,
            ]
            for column, value in enumerate(values, start=1):
                sheet.cell(row=current_row, column=column, value=value)
            sheet.cell(row=current_row, column=1).number_format = "dd/mm/yyyy hh:mm:ss"
            if interface.status.lower() != "up" or interface.protocol.lower() != "up":
                _set_excel_status_style(sheet.cell(row=current_row, column=6), "Critical")
                _set_excel_status_style(sheet.cell(row=current_row, column=7), "Critical")

    _style_excel_header(sheet, header_row, len(headers))
    _style_excel_body(sheet, header_row + 1, current_row, len(headers))
    if current_row > header_row:
        _add_excel_table(sheet, "LatestInterfacesTable", header_row, current_row, len(headers))
    else:
        sheet.cell(row=header_row + 1, column=1, value="Không có dữ liệu cổng mạng.")
    _set_excel_widths(sheet, [21, 16, 18, 24, 18, 13, 13, 14, 13, 15, 15, 14])
    sheet.freeze_panes = "A6"
    _configure_excel_print(sheet, landscape=True, repeat_rows="5:5")


def _build_findings_sheet(sheet, rows: list[ReportRow]) -> None:
    from openpyxl.styles import Alignment

    _prepare_detail_sheet(
        sheet,
        "CẢNH BÁO VÀ PHÁT HIỆN",
        "Danh sách phát hiện theo từng lần kiểm tra, đã chuyển nội dung sang tiếng Việt.",
        "DC2626",
    )
    headers = [
        "Thời gian kiểm tra",
        "Thiết bị",
        "IP quản trị",
        "Mức độ",
        "Chỉ số",
        "Nội dung",
        "Mã lần quét",
    ]
    header_row = 5
    for column, header in enumerate(headers, start=1):
        sheet.cell(row=header_row, column=column, value=header)
    current_row = header_row
    for device in rows:
        for check in device.checks:
            for finding in check.findings:
                current_row += 1
                values = [
                    _excel_datetime(check.checked_at),
                    _safe_excel_text(device.device_name),
                    device.host,
                    _status_label(finding.severity),
                    _metric_label(finding.metric),
                    _safe_excel_text(_translate_finding_message(finding.message)),
                    check.check_id,
                ]
                for column, value in enumerate(values, start=1):
                    sheet.cell(row=current_row, column=column, value=value)
                sheet.cell(row=current_row, column=1).number_format = "dd/mm/yyyy hh:mm:ss"
                _set_excel_status_style(sheet.cell(row=current_row, column=4), finding.severity)
                sheet.cell(row=current_row, column=6).alignment = Alignment(
                    horizontal="left", vertical="top", wrap_text=True
                )

    _style_excel_header(sheet, header_row, len(headers))
    _style_excel_body(sheet, header_row + 1, current_row, len(headers))
    if current_row > header_row:
        _add_excel_table(sheet, "HealthFindingsTable", header_row, current_row, len(headers))
    else:
        sheet.cell(
            row=header_row + 1,
            column=1,
            value="Không ghi nhận cảnh báo hoặc sự cố trong phạm vi báo cáo.",
        )
    _set_excel_widths(sheet, [21, 16, 18, 17, 23, 70, 14])
    sheet.freeze_panes = "A6"
    _configure_excel_print(sheet, landscape=True, repeat_rows="5:5")


def _prepare_excel_sheet(sheet, tab_color: str) -> None:
    from openpyxl.styles import Alignment, Font

    sheet.sheet_view.showGridLines = False
    sheet.sheet_properties.tabColor = tab_color
    sheet.sheet_view.zoomScale = 90
    sheet.sheet_format.defaultRowHeight = 18
    for row in sheet.iter_rows():
        for cell in row:
            cell.font = Font(name="Arial", size=10, color="334155")
            cell.alignment = Alignment(vertical="center")


def _prepare_detail_sheet(sheet, title: str, note: str, tab_color: str) -> None:
    from openpyxl.styles import Font

    _prepare_excel_sheet(sheet, tab_color)
    sheet["A2"] = title
    sheet["A2"].font = Font(name="Arial", size=15, bold=True, color="153E75")
    sheet["A3"] = note
    sheet["A3"].font = Font(name="Arial", size=10, italic=True, color="64748B")


def _style_excel_header(sheet, row: int, columns: int) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill, Side, Border

    fill = PatternFill("solid", fgColor="1F4E78")
    border = Border(bottom=Side(style="thin", color="FFFFFF"))
    for column in range(1, columns + 1):
        cell = sheet.cell(row=row, column=column)
        cell.fill = fill
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    sheet.row_dimensions[row].height = 34


def _style_excel_body(sheet, start_row: int, end_row: int, columns: int) -> None:
    from copy import copy

    from openpyxl.styles import Alignment, Font

    if end_row < start_row:
        return
    for row in sheet.iter_rows(
        min_row=start_row,
        max_row=end_row,
        min_col=1,
        max_col=columns,
    ):
        for cell in row:
            if not cell.font.bold:
                cell.font = Font(name="Arial", size=10, color="334155")
            alignment = copy(cell.alignment)
            alignment.vertical = alignment.vertical or "center"
            cell.alignment = alignment


def _add_excel_table(sheet, name: str, start_row: int, end_row: int, end_column: int) -> None:
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table, TableStyleInfo

    reference = f"A{start_row}:{get_column_letter(end_column)}{end_row}"
    table = Table(displayName=name, ref=reference)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    sheet.add_table(table)


def _set_excel_widths(sheet, widths: list[float]) -> None:
    from openpyxl.utils import get_column_letter

    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width


def _set_excel_status_style(cell, status: str) -> None:  # noqa: ANN001
    from openpyxl.styles import Font, PatternFill

    palette = {
        "Healthy": ("DCFCE7", "047857"),
        "Warning": ("FEF3C7", "B45309"),
        "Critical": ("FEE2E2", "B91C1C"),
        "Unknown": ("E2E8F0", "475569"),
    }
    fill, color = palette.get(status, palette["Unknown"])
    cell.fill = PatternFill("solid", fgColor=fill)
    cell.font = Font(name="Arial", size=10, bold=True, color=color)


def _set_excel_ssh_style(cell, ssh_ok: bool | None) -> None:  # noqa: ANN001
    if ssh_ok is True:
        _set_excel_status_style(cell, "Healthy")
    elif ssh_ok is False:
        _set_excel_status_style(cell, "Critical")
    else:
        _set_excel_status_style(cell, "Unknown")


def _configure_excel_print(sheet, *, landscape: bool, repeat_rows: str) -> None:
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape" if landscape else "portrait"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_title_rows = repeat_rows
    sheet.sheet_properties.outlinePr.summaryBelow = True
    sheet.oddFooter.center.text = "Network Health Check"
    sheet.oddFooter.right.text = "Trang &P / &N"


def _excel_datetime(value: datetime | None) -> datetime | None:
    if value is None:
        return None
    if value.tzinfo is None:
        value = value.replace(tzinfo=timezone.utc)
    return value.astimezone(VIETNAM_TZ).replace(tzinfo=None)


def _excel_percent(value: float | None) -> float | None:
    return None if value is None else value / 100


def _safe_excel_text(value: str) -> str:
    if value in {"", "-"}:
        return value
    return f"'{value}" if value.startswith(("=", "+", "-", "@")) else value


def _check_row(check: HealthCheck) -> ReportCheckRow:
    preferred_ping = next(
        (ping for ping in check.pings if ping.source_type == "device"),
        None,
    )
    if preferred_ping is None:
        preferred_ping = next(
            (ping for ping in check.pings if ping.source_type == "backend"),
            None,
        )
    return ReportCheckRow(
        check_id=check.id,
        checked_at=check.checked_at or check.finished_at or check.started_at,
        status=check.overall_status or "Unknown",
        ssh_ok=check.ssh_ok,
        uptime_seconds=check.uptime_seconds,
        cpu_percent=check.cpu_percent,
        ram_percent=check.ram_percent,
        ping_target=preferred_ping.target if preferred_ping else "-",
        packet_loss_percent=(
            preferred_ping.packet_loss_percent if preferred_ping else None
        ),
        latency_avg_ms=preferred_ping.latency_avg_ms if preferred_ping else None,
        error=check.error or "",
        interfaces=tuple(
            ReportInterfaceRow(
                name=item.name,
                ip_address=item.ip_address,
                status=item.status,
                protocol=item.protocol,
                input_errors=item.input_errors,
                crc_errors=item.crc_errors,
                output_errors=item.output_errors,
                output_drops=item.output_drops,
            )
            for item in check.interfaces
        ),
        findings=tuple(
            ReportFindingRow(
                severity=item.severity,
                metric=item.metric,
                message=item.message,
            )
            for item in check.findings
        ),
    )


def _register_pdf_fonts() -> tuple[str, str]:
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    candidates = (
        (
            "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
            "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
        ),
        ("C:/Windows/Fonts/arial.ttf", "C:/Windows/Fonts/arialbd.ttf"),
    )
    for regular_path, bold_path in candidates:
        if Path(regular_path).is_file() and Path(bold_path).is_file():
            if "ReportFont" not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(TTFont("ReportFont", regular_path))
            if "ReportFontBold" not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(TTFont("ReportFontBold", bold_path))
            return "ReportFont", "ReportFontBold"
    return "Helvetica", "Helvetica-Bold"


def _kpi_table(
    items: list[tuple[str, object, str]],
    font_name: str,
    bold_font_name: str,
):
    from reportlab.lib import colors
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, Table, TableStyle

    label_style = ParagraphStyle(
        "KpiLabel",
        fontName=font_name,
        fontSize=7,
        leading=9,
        alignment=1,
        textColor=colors.HexColor("#64748B"),
    )
    value_styles = [
        ParagraphStyle(
            f"KpiValue{index}",
            fontName=bold_font_name,
            fontSize=15,
            leading=18,
            alignment=1,
            textColor=colors.HexColor(color),
        )
        for index, (_, _, color) in enumerate(items)
    ]
    table = Table(
        [
            [Paragraph(label, label_style) for label, _, _ in items],
            [
                Paragraph(str(value), value_styles[index])
                for index, (_, value, _) in enumerate(items)
            ],
        ],
        colWidths=[44 * mm for _ in items],
    )
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F8FAFC")),
                ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
                ("INNERGRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#E2E8F0")),
                ("TOPPADDING", (0, 0), (-1, -1), 5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ]
        )
    )
    return table


def _report_table_style(font_name: str):
    from reportlab.lib import colors
    from reportlab.platypus import TableStyle

    return TableStyle(
        [
            ("FONTNAME", (0, 0), (-1, -1), font_name),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E78")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#CBD5E1")),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("ALIGN", (0, 0), (-1, 0), "CENTER"),
            ("FONTSIZE", (0, 0), (-1, -1), 7),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
            ("LEFTPADDING", (0, 0), (-1, -1), 3),
            ("RIGHTPADDING", (0, 0), (-1, -1), 3),
            ("TOPPADDING", (0, 0), (-1, -1), 3.5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3.5),
        ]
    )


def _apply_status_colors(table, rows, *, status_column: int) -> None:  # noqa: ANN001
    from reportlab.lib import colors

    palette = {
        "Healthy": colors.HexColor("#047857"),
        "Warning": colors.HexColor("#B45309"),
        "Critical": colors.HexColor("#B91C1C"),
        "Unknown": colors.HexColor("#64748B"),
    }
    statuses = [
        item.latest_status if isinstance(item, ReportRow) else str(item)
        for item in rows
    ]
    for row_index, status in enumerate(statuses, start=1):
        table.setStyle(
            [("TEXTCOLOR", (status_column, row_index), (status_column, row_index), palette.get(status, palette["Unknown"]))]
        )


def _report_period(
    checks: list[ReportCheckRow],
    date_from: datetime | None,
    date_to: datetime | None,
) -> tuple[datetime | None, datetime | None]:
    values = [check.checked_at for check in checks]
    return date_from or (min(values) if values else None), date_to or (max(values) if values else None)


def _average(values: list[float]) -> float | None:
    return round(fmean(values), 2) if values else None


def _display(value: float | None) -> str:
    return "-" if value is None else f"{value:.2f}"


def _percent(value: float | None) -> str:
    return "-" if value is None else f"{value:.1f}%"


def _milliseconds(value: float | None) -> str:
    return "-" if value is None else f"{value:.1f} ms"


def _format_datetime(value: datetime | None) -> str:
    if value is None:
        return "-"
    if value.tzinfo is None:
        value = value.replace(tzinfo=timezone.utc)
    return value.astimezone(VIETNAM_TZ).strftime("%H:%M:%S %d/%m/%Y")


def _period_label(date_from: datetime | None, date_to: datetime | None) -> str:
    if date_from is None and date_to is None:
        return "Không có dữ liệu"
    return f"{_format_datetime(date_from)} đến {_format_datetime(date_to)}"


def _format_uptime(seconds: int | None) -> str:
    if seconds is None:
        return "-"
    days, remainder = divmod(seconds, 86_400)
    hours, remainder = divmod(remainder, 3_600)
    minutes = remainder // 60
    parts = []
    if days:
        parts.append(f"{days} ngày")
    if hours:
        parts.append(f"{hours} giờ")
    if minutes or not parts:
        parts.append(f"{minutes} phút")
    return " ".join(parts)


def _status_label(value: str) -> str:
    return {
        "Healthy": "Tốt",
        "Warning": "Cảnh báo",
        "Critical": "Nghiêm trọng",
        "Unknown": "Chưa xác định",
    }.get(value, value)


def _ssh_label(value: bool | None) -> str:
    if value is True:
        return "Thành công"
    if value is False:
        return "Thất bại"
    return "Chưa rõ"


def _metric_label(value: str) -> str:
    return {
        "cpu": "CPU",
        "ram": "RAM",
        "ssh": "SSH",
        "management_ping": "Ping quản trị",
        "path_ping": "Ping đường truyền",
        "packet_loss": "Mất gói",
        "latency": "Độ trễ",
        "interface": "Cổng mạng",
        "interface_errors": "Lỗi cổng mạng",
        "collector": "Thu thập dữ liệu",
    }.get(value.lower(), value)


def _translate_finding_message(message: str) -> str:
    exact = {
        "SSH connection failed": "Kết nối SSH thất bại",
        "Management ping data is unavailable": "Không có dữ liệu ping quản trị",
        "Path ping data is unavailable": "Không có dữ liệu ping đường truyền",
        "Path latency data is unavailable": "Không có dữ liệu độ trễ đường truyền",
    }
    if message in exact:
        return exact[message]
    if message.lower().startswith("collection failed"):
        return "Thu thập dữ liệu thất bại: thiết bị không phản hồi qua SSH."
    patterns = (
        (r"^CPU usage is (.+)$", r"Mức sử dụng CPU là \1"),
        (r"^RAM usage is (.+)$", r"Mức sử dụng RAM là \1"),
        (r"^CPU data is unavailable$", "Không có dữ liệu CPU"),
        (r"^RAM data is unavailable$", "Không có dữ liệu RAM"),
        (r"^Management packet loss is (.+)$", r"Tỷ lệ mất gói quản trị là \1"),
        (r"^Path packet loss is (.+)$", r"Tỷ lệ mất gói đường truyền là \1"),
        (r"^Average latency is (.+)$", r"Độ trễ trung bình là \1"),
        (r"^(.+) is (.+)/(.+)$", r"Cổng \1 có trạng thái \2/\3"),
        (r"^(.+) has non-zero error counters$", r"Cổng \1 có bộ đếm lỗi khác 0"),
    )
    for pattern, replacement in patterns:
        if re.match(pattern, message, flags=re.IGNORECASE):
            return re.sub(pattern, replacement, message, flags=re.IGNORECASE)
    return message


def _translate_error(message: str) -> str:
    if not message:
        return "-"
    lowered = message.lower()
    if "tcp connection to device failed" in lowered:
        return "Không thể kết nối TCP/SSH. Kiểm tra IP, cổng 22, firewall và trạng thái GNS3."
    if "authentication failed" in lowered:
        return "Xác thực SSH thất bại. Kiểm tra tên đăng nhập và mật khẩu."
    if "timed out" in lowered or "timeout" in lowered:
        return "Kết nối đến thiết bị quá thời gian chờ."
    return message


def _cell(value: object, style):  # noqa: ANN001
    from reportlab.platypus import Paragraph

    safe_value = str(value).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    return Paragraph(safe_value, style)


def _safe_csv(value: str) -> str:
    if value in {"", "-"}:
        return value
    return f"'{value}" if value.startswith(("=", "+", "-", "@")) else value
